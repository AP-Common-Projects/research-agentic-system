"""v3 rich Excel deliverable — the client's explicit ask: a workbook they can
filter, sort, and search themselves (not an AI-written report). Covers the
sanitization helpers, the AutoFilter/formatting contract on every sheet, and
that export_excel assembles the right sheets from the store without touching
a LangGraph checkpoint.
"""

from __future__ import annotations

from datetime import datetime, timezone
from unittest.mock import MagicMock, patch

from src.export import (
    _excel_safe,
    _write_excel_sheet,
    build_excel_workbook_v3,
    export_excel,
)


class TestExcelSafe:
    def test_strips_tz_aware_datetime(self):
        dt = datetime(2026, 1, 1, tzinfo=timezone.utc)
        result = _excel_safe(dt)
        assert result.tzinfo is None
        assert result == datetime(2026, 1, 1)

    def test_naive_datetime_passes_through(self):
        dt = datetime(2026, 1, 1)
        assert _excel_safe(dt) == dt

    def test_strips_illegal_control_characters(self):
        assert _excel_safe("hello\x00world\x01") == "helloworld"

    def test_joins_lists_for_a_flat_cell(self):
        assert _excel_safe(["country_code", "region"]) == "country_code; region"

    def test_passes_through_plain_values(self):
        assert _excel_safe(42) is 42 or _excel_safe(42) == 42
        assert _excel_safe(None) is None
        assert _excel_safe(True) is True


class TestWriteExcelSheet:
    def test_empty_rows_writes_placeholder(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        _write_excel_sheet(ws, [])
        assert ws["A1"].value == "(no rows)"

    def test_headers_bold_and_frozen_and_filterable(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        rows = [
            {"channel_id": "c1", "subscriber_count": 100000},
            {"channel_id": "c2", "subscriber_count": 250000},
        ]
        _write_excel_sheet(ws, rows)

        assert [c.value for c in ws[1]] == ["channel_id", "subscriber_count"]
        assert all(c.font.bold for c in ws[1])
        assert ws.freeze_panes == "A2"
        # AutoFilter must cover the full used range, header through last row —
        # this is the literal "can filter the data columns" requirement.
        assert ws.auto_filter.ref == "A1:B3"
        assert ws["A2"].value == "c1"
        assert ws["B3"].value == 250000

    def test_sanitizes_cell_values_on_write(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        rows = [{"description": "spam\x00text", "missing_required_fields": ["a", "b"]}]
        _write_excel_sheet(ws, rows)
        assert ws["A2"].value == "spamtext"
        assert ws["B2"].value == "a; b"


class TestBuildExcelWorkbookV3:
    def _manifest(self):
        return {
            "niche": "crime",
            "channels": 2,
            "videos": 3,
            "niches_covered": 1,
            "total_cost_usd": 1.23,
            "exported_at": datetime.now(timezone.utc).isoformat(),
        }

    def test_produces_all_seven_sheets(self):
        wb = build_excel_workbook_v3(
            self._manifest(),
            channels=[{"channel_id": "c1", "title": "X"}],
            videos=[{"video_id": "v1", "title": "Y"}],
            niches=[{"category": "crime", "sub_niche": "true_crime_documentary", "channel_count": 1}],
            success_factors=[{"channel_id": "c1", "factor_code": "consistent_upload_cadence"}],
            failure_factors=[],
        )
        assert wb.sheetnames == [
            "Overview", "Channels", "Videos", "Shorts", "Niches",
            "Success Factors", "Failure Factors",
        ]

    def test_overview_carries_run_stats_and_niche_rollup(self):
        wb = build_excel_workbook_v3(
            self._manifest(),
            channels=[{"channel_id": "c1"}, {"channel_id": "c2"}],
            videos=[],
            niches=[{"category": "crime", "sub_niche": "true_crime_documentary", "channel_count": 2}],
            success_factors=[],
            failure_factors=[],
        )
        ws = wb["Overview"]
        # iter_rows pads every row out to the sheet's widest row, so compare
        # only the leading cells each row actually carries.
        rows = [tuple(c.value for c in row[:3]) for row in ws.iter_rows()]
        assert ("Channels", 2, None) in rows
        assert ("crime", "true_crime_documentary", 2) in rows

    def test_empty_failure_factors_sheet_does_not_crash(self):
        wb = build_excel_workbook_v3(
            self._manifest(), channels=[], videos=[], niches=[],
            success_factors=[], failure_factors=[],
        )
        assert wb["Failure Factors"]["A1"].value == "(no rows)"


class TestExportExcel:
    def test_assembles_sheets_from_the_store_not_a_checkpoint(self, tmp_path):
        """export_excel must not touch _load_checkpoint / final_report at
        all — it is the v3 dataset-first deliverable, independent of
        whether a run's LangGraph thread checkpoint still exists."""
        out_path = tmp_path / "crime.xlsx"
        with (
            patch("src.export.fetch_run_channels", return_value=[{"channel_id": "c1", "subscriber_count": 60000}]),
            patch("src.export.fetch_run_videos", return_value=[{"video_id": "v1", "outlier_score": 3.2}]),
            patch(
                "src.export.fetch_run_niche_breakdown",
                return_value=[{"category": "crime", "sub_niche": "true_crime_documentary", "channel_count": 1}],
            ),
            patch("src.export.fetch_run_success_factors", return_value=[]),
            patch("src.export.fetch_run_failure_factors", return_value=[]),
            patch("src.export.dominant_run_category", return_value="crime"),
            patch("src.export._fetch", return_value=[{"total_cost_usd": 4.5}]),
        ):
            result = export_excel("run-crime-quick", out_path)

        assert result == out_path
        assert out_path.exists()

        from openpyxl import load_workbook

        wb = load_workbook(out_path)
        assert wb["Channels"]["A2"].value == "c1"
        assert wb["Overview"]["A3"].value == "Channels"
        assert wb["Overview"]["B3"].value == 1

    def test_scopes_every_sheet_to_the_runs_dominant_category(self, tmp_path):
        """The client's two exclusion rules — over-50k-subs only, and no
        unrelated channels — must reach EVERY sheet, not just Channels. A
        crime-seeded run legitimately discovers a lifestyle vlog named
        "True Crime PROFILE 2026"; its videos and factors must not leak
        into the deliverable through a sheet that forgot to filter."""
        out_path = tmp_path / "crime.xlsx"
        with (
            patch("src.export.fetch_run_channels", return_value=[]) as ch,
            patch("src.export.fetch_run_videos", return_value=[]) as vid,
            patch("src.export.fetch_run_niche_breakdown", return_value=[]) as nic,
            patch("src.export.fetch_run_success_factors", return_value=[]) as suc,
            patch("src.export.fetch_run_failure_factors", return_value=[]) as fail,
            patch("src.export.dominant_run_category", return_value="crime"),
            patch("src.export._fetch", return_value=[]),
        ):
            export_excel("run-crime-quick", out_path)

        assert ch.call_args.args[1] == "crime"
        assert vid.call_args.args[2] == "crime"
        for mock in (nic, suc, fail):
            assert mock.call_args.args[1] == "crime"

    def test_all_categories_keeps_every_discovered_channel(self, tmp_path):
        out_path = tmp_path / "all.xlsx"
        with (
            patch("src.export.fetch_run_channels", return_value=[]) as ch,
            patch("src.export.fetch_run_videos", return_value=[]),
            patch("src.export.fetch_run_niche_breakdown", return_value=[]),
            patch("src.export.fetch_run_success_factors", return_value=[]),
            patch("src.export.fetch_run_failure_factors", return_value=[]),
            patch("src.export.dominant_run_category") as dominant,
            patch("src.export._fetch", return_value=[]),
        ):
            export_excel("run-x", out_path, all_categories=True)

        assert ch.call_args.args[1] is None
        dominant.assert_not_called()

    def test_all_channels_drops_the_subscriber_floor_and_video_cap(self, tmp_path):
        """The big-run deliverable's own ask: every discovered channel, no
        50k floor, and every tagged video with no export_max_videos cap.
        Implies all_categories too — a sub-floor channel let back in by
        min_subscribers=0 must not then get dropped by the category filter
        instead."""
        out_path = tmp_path / "all_channels.xlsx"
        with (
            patch("src.export.fetch_run_channels", return_value=[]) as ch,
            patch("src.export.fetch_run_videos", return_value=[]) as vid,
            patch("src.export.fetch_run_niche_breakdown", return_value=[]) as nic,
            patch("src.export.fetch_run_success_factors", return_value=[]) as suc,
            patch("src.export.fetch_run_failure_factors", return_value=[]) as fail,
            patch("src.export.dominant_run_category") as dominant,
            patch("src.export._fetch", return_value=[]),
        ):
            export_excel("run-x", out_path, all_channels=True)

        assert ch.call_args.args[1] is None, "all_channels must imply all_categories"
        assert ch.call_args.args[2] == 0, "min_subscribers=0 — no floor"
        assert vid.call_args.args[1] is None, "no export_max_videos cap"
        assert vid.call_args.args[3] == 0
        for mock in (nic, suc, fail):
            assert mock.call_args.args[2] == 0
        dominant.assert_not_called()

    def test_video_cap_and_channel_scope_are_independent(self, tmp_path):
        """The 50k floor and category filter are the client's scope
        requirement and stay on by default. The Videos row cap is a
        SEPARATE, independent knob (cap_videos) — dropping the cap must
        not silently drop the floor/category scope, and keeping the scope
        must not silently reimpose the cap. Regression for conflating the
        two: an earlier version only offered all_channels, which coupled
        'no video cap' to 'no floor, no category filter' as one flag."""
        out_path = tmp_path / "scoped_uncapped.xlsx"
        with (
            patch("src.export.fetch_run_channels", return_value=[]) as ch,
            patch("src.export.fetch_run_videos", return_value=[]) as vid,
            patch("src.export.fetch_run_niche_breakdown", return_value=[]),
            patch("src.export.fetch_run_success_factors", return_value=[]),
            patch("src.export.fetch_run_failure_factors", return_value=[]),
            patch("src.export.dominant_run_category", return_value="finance"),
            patch("src.export._fetch", return_value=[]),
        ):
            export_excel("run-x", out_path)  # every default: no flags passed

        assert ch.call_args.args[1] == "finance", "category scope must stay on by default"
        assert ch.call_args.args[2] is None, "50k floor (via _floor()) must stay on by default"
        assert vid.call_args.args[1] is None, "no row cap by default — cap_videos defaults False"
        assert vid.call_args.args[2] == "finance", "videos share the channels' category scope"

    def test_defaults_output_path_under_export_dir(self, tmp_path):
        with (
            patch("src.export.export_dir", return_value=tmp_path),
            patch("src.export.fetch_run_channels", return_value=[]),
            patch("src.export.fetch_run_videos", return_value=[]),
            patch("src.export.fetch_run_niche_breakdown", return_value=[]),
            patch("src.export.fetch_run_success_factors", return_value=[]),
            patch("src.export.fetch_run_failure_factors", return_value=[]),
            patch("src.export.dominant_run_category", return_value=None),
            patch("src.export._fetch", return_value=[]),
        ):
            result = export_excel("run-x")

        assert result == tmp_path / "run-x.xlsx"
        assert result.exists()


class TestSplitShortsAndFormatting:
    """The client brief is explicit: "Keep Shorts separate. Do not mix Shorts
    with long-form analysis." That makes the split a correctness requirement
    of the deliverable, not a presentation preference."""

    def test_shorts_and_long_form_go_to_different_sheets(self):
        from src.export import _split_shorts

        long_form, shorts = _split_shorts([
            {"video_id": "a", "is_short": False, "duration_seconds": 600},
            {"video_id": "b", "is_short": True, "duration_seconds": 45},
        ])
        assert [v["video_id"] for v in long_form] == ["a"]
        assert [v["video_id"] for v in shorts] == ["b"]

    def test_each_sheet_reports_duration_in_its_own_unit(self):
        """Minutes for long-form, raw seconds for Shorts — a Short shown in
        minutes is a fraction on every row, and a two-hour documentary shown
        in seconds is unreadable."""
        from src.export import _split_shorts

        long_form, shorts = _split_shorts([
            {"video_id": "a", "is_short": False, "duration_seconds": 600},
            {"video_id": "b", "is_short": True, "duration_seconds": 45},
        ])
        assert long_form[0]["duration(in minute -by default)"] == 10.0
        assert "duration_seconds" not in long_form[0]
        assert shorts[0]["duration_seconds"] == 45
        assert "duration(in minute -by default)" not in shorts[0]

    def test_long_form_past_an_hour_switches_to_hours_and_minutes(self):
        from src.export import _split_shorts

        long_form, _ = _split_shorts(
            [{"video_id": "a", "is_short": False, "duration_seconds": 5400}]
        )
        assert long_form[0]["duration(in minute -by default)"] == "1h 30m"

    def test_is_short_column_is_dropped_from_both_sheets(self):
        """Sheet membership already carries it; a redundant column is one
        more thing that can disagree with the sheet it sits in."""
        from src.export import _split_shorts

        long_form, shorts = _split_shorts([
            {"video_id": "a", "is_short": False, "duration_seconds": 600},
            {"video_id": "b", "is_short": True, "duration_seconds": 45},
        ])
        assert "is_short" not in long_form[0]
        assert "is_short" not in shorts[0]

    def test_numeric_columns_get_thousands_separators(self):
        wb = build_excel_workbook_v3(
            _MANIFEST_FOR_FORMATTING,
            channels=[{"channel_id": "c1", "subscriber_count": 13100000}],
            videos=[], niches=[], success_factors=[], failure_factors=[],
        )
        ws = wb["Channels"]
        assert ws.cell(row=2, column=2).number_format == "#,##0"

    def test_booleans_are_not_comma_formatted_as_numbers(self):
        """Booleans are ints in Python — formatting one as a number would
        render TRUE as "1"."""
        wb = build_excel_workbook_v3(
            _MANIFEST_FOR_FORMATTING,
            channels=[{"channel_id": "c1", "is_likely_news": True}],
            videos=[], niches=[], success_factors=[], failure_factors=[],
        )
        ws = wb["Channels"]
        assert ws.cell(row=2, column=2).number_format != "#,##0"


_MANIFEST_FOR_FORMATTING = {
    "run_id": "run-x", "niche": "crime", "channels": 1, "videos": 0,
    "niches_covered": 0, "total_cost_usd": 0, "exported_at": "2026-08-30T00:00:00+00:00",
}
