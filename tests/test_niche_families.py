"""Every niche gets a family, and every family carries at least ten.

A Standard education run shipped 135 sub-niches for 167 channels, 115 of
them holding exactly one channel. Individually correct, collectively
unreadable: there is no comparison to be made between 115 groups of one.

The tier that fixes it already existed and was never filled in. The
`primary_niche_groups` table has carried a Primary_Niche level since the
crime brief and the Channels sheet has read it since, but nine groups were
hand-seeded for crime and finance and nothing assigned one for any other
vertical -- 618 of 2,219 niches had a family. The completeness gate said so
in as many words: "a real limitation, but not one a run can close".
"""

from __future__ import annotations

import pathlib

import pytest

from src.nodes.assign_niche_families import (
    MIN_SUB_NICHES_PER_FAMILY,
    _enforce_minimum,
    _nearest_family,
    _split_kept_and_loose,
    _target_family_count,
)


def _niches(spec):
    """(name, category, family) triples -> the node's niche dicts."""
    return [
        {"niche_id": i, "niche_name": name, "parent_category": cat,
         "description": "", "current_family": fam, "channel_count": 1}
        for i, (name, cat, fam) in enumerate(spec)
    ]


class TestTheMinimumIsEnforcedOnTheOutput:
    """A model told "at least ten" still returns a family of one."""

    def test_no_family_survives_under_the_minimum(self):
        niches = _niches([(f"n{i}", "education", None) for i in range(40)])
        placed = {n["niche_id"]: f"F{n['niche_id'] % 8}" for n in niches}
        out = _enforce_minimum(placed, niches, [])
        sizes = {}
        for label in out.values():
            sizes[label] = sizes.get(label, 0) + 1
        assert sizes, out
        assert min(sizes.values()) >= MIN_SUB_NICHES_PER_FAMILY, sizes

    def test_every_niche_keeps_a_family(self):
        niches = _niches([(f"n{i}", "education", None) for i in range(40)])
        placed = {n["niche_id"]: f"F{n['niche_id'] % 8}" for n in niches}
        out = _enforce_minimum(placed, niches, [])
        assert set(out) == {n["niche_id"] for n in niches}
        assert all(out.values())

    def test_a_niche_the_model_skipped_is_not_dropped(self):
        niches = _niches([(f"n{i}", "education", None) for i in range(30)])
        placed = {n["niche_id"]: "Big" for n in niches[:25]}
        out = _enforce_minimum(placed, niches, [{"label": "Big"}])
        assert len(out) == 30

    def test_a_run_too_small_for_one_family_gets_exactly_one(self):
        """The rule is unsatisfiable below ten niches. One honest family
        beats three of two."""
        niches = _niches([(f"n{i}", "education", None) for i in range(6)])
        placed = {n["niche_id"]: f"F{n['niche_id'] % 3}" for n in niches}
        out = _enforce_minimum(placed, niches, [])
        assert len(set(out.values())) == 1

    def test_it_terminates_when_nothing_can_reach_the_minimum(self):
        niches = _niches([(f"n{i}", "education", None) for i in range(11)])
        placed = {n["niche_id"]: f"F{n['niche_id']}" for n in niches}
        out = _enforce_minimum(placed, niches, [])
        assert len(set(out.values())) == 1


class TestMergesFollowMeaningNotSize:
    """Merging into whichever family was biggest filed "Technology &
    Programming Education" under "Lifestyle & Family Vlogging"."""

    def test_the_nearest_family_by_category_wins_over_the_largest(self):
        profiles = {
            "Doomed": {"technology": 1.0},
            "Big But Unrelated": {"lifestyle": 1.0},
            "Small And Related": {"technology": 1.0},
        }
        counts = {"Doomed": 2, "Big But Unrelated": 40, "Small And Related": 12}
        assert _nearest_family(
            profiles["Doomed"], ["Big But Unrelated", "Small And Related"],
            profiles, counts,
        ) == "Small And Related"

    def test_size_breaks_a_tie_so_the_merge_terminates(self):
        profiles = {"Doomed": {"a": 1.0}, "X": {"b": 1.0}, "Y": {"b": 1.0}}
        counts = {"Doomed": 1, "X": 5, "Y": 9}
        assert _nearest_family(profiles["Doomed"], ["X", "Y"], profiles, counts) == "Y"

    def test_a_real_merge_lands_somewhere_sensible(self):
        spec = ([(f"edu{i}", "education", None) for i in range(12)]
                + [(f"life{i}", "lifestyle", None) for i in range(12)]
                + [(f"sci{i}", "science_explainer", None) for i in range(3)])
        niches = _niches(spec)
        placed = {}
        for n in niches:
            placed[n["niche_id"]] = {
                "education": "Academic", "lifestyle": "Living",
                "science_explainer": "Academic",
            }[n["parent_category"]] if n["parent_category"] != "science_explainer" else "Science"
        out = _enforce_minimum(placed, niches, [])
        science_went_to = {out[n["niche_id"]] for n in niches
                           if n["parent_category"] == "science_explainer"}
        assert science_went_to == {"Academic"}, science_went_to


class TestFamiliesAlreadyBigEnoughAreLeftAlone:
    def test_a_curated_group_with_enough_members_survives(self):
        spec = [(f"c{i}", "crime", "Police Investigation Documentary")
                for i in range(14)]
        kept, loose = _split_kept_and_loose(_niches(spec))
        assert len(kept) == 14
        assert loose == []

    def test_a_curated_group_of_one_is_re_filed(self):
        """Three stray finance niches rendered as three families of one on
        an education workbook. That is the failure, curated or not."""
        spec = ([("fin1", "finance", "Personal Finance")]
                + [(f"e{i}", "education", None) for i in range(12)])
        kept, loose = _split_kept_and_loose(_niches(spec))
        assert kept == {}
        assert len(loose) == 13

    def test_an_ungrouped_niche_is_always_loose(self):
        kept, loose = _split_kept_and_loose(_niches([("a", "education", None)]))
        assert kept == {} and len(loose) == 1


class TestTheTargetCountCannotBreakItsOwnRule:
    @pytest.mark.parametrize("n,expected", [(135, 13), (100, 10), (9, 1), (0, 1)])
    def test_integer_division_keeps_the_promise(self, n, expected):
        assert _target_family_count(n) == expected

    def test_the_target_never_forces_a_short_family(self):
        for n in range(1, 600):
            assert _target_family_count(n) * MIN_SUB_NICHES_PER_FAMILY <= max(n, MIN_SUB_NICHES_PER_FAMILY)


class TestTheWorkbookCarriesIt:
    def test_the_families_sheet_is_one_of_the_swept_sheets(self):
        from src.export import sheet_rows

        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        body = src[src.index("def sheet_rows("):src.index("def export_excel(")]
        assert '"Niche Families"' in body

    def test_the_niches_sheet_names_each_niches_family(self):
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        body = src[src.index("def fetch_run_niche_breakdown("):
                   src.index("def fetch_run_niche_families(")]
        assert "niche_family" in body

    def test_the_workbook_writes_the_sheet(self):
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        assert 'wb.create_sheet("Niche Families")' in src

    def test_the_gate_no_longer_calls_primary_niche_unclosable(self):
        src = pathlib.Path("src/tools/export_completeness.py").read_text(encoding="utf-8")
        assert "not one a run can close" not in src

    def test_the_node_is_in_the_graph_before_cohorts(self):
        src = pathlib.Path("src/graph.py").read_text(encoding="utf-8")
        assert 'graph.add_edge("populate_shared_fields", "assign_niche_families")' in src
        assert 'graph.add_edge("assign_niche_families", "assign_cohorts")' in src

    def test_the_run_that_prompted_this_now_has_families(self):
        """run-e4e794436210: 135 sub-niches, 115 of them singletons."""
        from src.export import fetch_run_niche_families

        fams = fetch_run_niche_families("run-e4e794436210", None, None, True)
        assert fams, "no families assigned"
        for f in fams:
            assert int(f["distinct_sub_niches"]) >= MIN_SUB_NICHES_PER_FAMILY, f


class TestTheOverviewSummaryShowsFamilies:
    """The Overview sheet had a table headed "Niche family" that listed all
    135 SUB-niches, one per row, 115 of them holding a single channel.

    It read `primary_topic`, `distinct_sub_niches`, `example_sub_niches` and
    `pct_of_vertical` -- four keys the niche breakdown has never produced --
    so three of its columns shipped blank and its label fell through to the
    sub-niche name. A heading that promises a rollup over a table that is
    not one.
    """

    def test_the_summary_is_built_from_the_family_rows(self):
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        body = src[src.index("def build_excel_workbook_v3("):
                   src.index('_write_excel_sheet(wb.create_sheet("Channels")')]
        assert 'ws.append(["Niche family", "Channels", "Share of run",' in body

    def test_it_no_longer_reads_keys_nothing_produces(self):
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        body = src[src.index("def build_excel_workbook_v3("):
                   src.index('_write_excel_sheet(wb.create_sheet("Channels")')]
        assert 'n.get("primary_topic")' not in body
        assert 'n.get("pct_of_vertical")' not in body

    def test_the_family_count_is_stated(self):
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        assert 'ws.append(["Niche families", len(families or ())])' in src

    def test_the_delivered_workbook_summarises_five_rows_not_135(self):
        from openpyxl import load_workbook

        path = pathlib.Path("exports/run-e4e794436210/run-e4e794436210.xlsx")
        if not path.exists():
            pytest.skip("the reference workbook has not been exported here")
        wb = load_workbook(path, read_only=True)
        rows = [r for r in wb["Overview"].values]
        wb.close()
        header = next(i for i, r in enumerate(rows) if r and r[0] == "Niche family")
        table = [r for r in rows[header + 1:] if r and r[0]]
        assert len(table) <= 13, len(table)
        for row in table:
            assert int(row[3]) >= MIN_SUB_NICHES_PER_FAMILY, row


class TestTheWrittenFileIsChecked:
    """Enforced upstream is what every completeness failure in this project
    was claiming when it shipped. The rule is about what the client opens."""

    def test_the_families_sheet_is_verified(self):
        from src.tools.workbook_verify import _DATA_SHEETS

        assert "Niche Families" in _DATA_SHEETS

    def test_a_short_family_on_the_sheet_fails_verification(self):
        from src.tools.workbook_verify import _short_families

        header = ["niche_family", "distinct_sub_niches"]
        body = [("Big", 40), ("Also Big", 12), ("Tiny", 3)]
        assert _short_families(header, body) == [("Tiny", 3)]

    def test_families_meeting_the_rule_pass(self):
        from src.tools.workbook_verify import _short_families

        assert _short_families(
            ["niche_family", "distinct_sub_niches"], [("A", 10), ("B", 40)]
        ) == []

    def test_a_workbook_too_small_for_one_family_is_exempt(self):
        """The rule is unsatisfiable below ten sub-niches; one honest
        family holding everything is the right answer, not a fault."""
        from src.tools.workbook_verify import _short_families

        assert _short_families(
            ["niche_family", "distinct_sub_niches"], [("Only", 6)]
        ) == []

    def test_an_older_workbook_without_the_columns_is_not_failed(self):
        from src.tools.workbook_verify import _short_families

        assert _short_families(["category", "sub_niche"], [("a", "b")]) == []

    def test_the_report_says_so_and_is_not_ok(self):
        from src.tools.workbook_verify import VerifyReport

        r = VerifyReport(path="x.xlsx", short_families=[("Tiny", 3)])
        assert not r.ok
        assert "SHORT FAMILY: Tiny holds 3 sub-niches, needs 10" in r.render()


class TestEveryViewOfARunAgrees:
    """Three surfaces were answering "what is a niche family" three
    different ways: primary_niche_groups in the workbook, parent_category in
    the export graph, and primary_topic in the console tree -- which showed
    138 families for a run whose workbook had five."""

    def test_the_console_tree_uses_the_family_tier(self):
        src = pathlib.Path("src/api/deliverables.py").read_text(encoding="utf-8")
        body = src[src.index("def workbook_tree("):]
        assert "primary_niche_groups png" in body
        assert "png.group_label" in body

    def test_the_console_tree_still_falls_back_for_older_workbooks(self):
        src = pathlib.Path("src/api/deliverables.py").read_text(encoding="utf-8")
        body = src[src.index("def workbook_tree("):]
        assert "NULLIF(c.primary_topic, '')" in body

    def test_the_export_graph_carries_the_family(self):
        from src.export import build_graph_payload

        nodes = build_graph_payload(
            [{"channel_id": "c1", "title": "One", "subscriber_count": 9,
              "discovery_method": "keyword", "primary_niche": "Exam Prep",
              "sub_niche": "maths"}],
            [],
        )["nodes"]
        assert nodes[0]["family"] == "Exam Prep"
        assert nodes[0]["sub_niche"] == "maths"

    def test_the_graph_only_offers_families_the_workbook_has(self):
        """A channel an earlier run found dragged "Personal Finance" into
        an education run's family filter, one node in it, beside a workbook
        promising ten sub-niches a family."""
        from src.export import build_graph_payload

        nodes = build_graph_payload(
            [{"channel_id": "c1", "title": "One", "subscriber_count": 9,
              "discovery_method": "keyword", "primary_niche": "Personal Finance"},
             {"channel_id": "c2", "title": "Two", "subscriber_count": 9,
              "discovery_method": "keyword", "primary_niche": "Exam Prep"}],
            [], families={"Exam Prep"},
        )["nodes"]
        by_id = {n["id"]: n for n in nodes}
        assert by_id["c1"]["family"] == ""
        assert by_id["c2"]["family"] == "Exam Prep"

    def test_no_vocabulary_given_means_no_filtering(self):
        from src.export import build_graph_payload

        nodes = build_graph_payload(
            [{"channel_id": "c1", "title": "One", "subscriber_count": 9,
              "discovery_method": "keyword", "primary_niche": "Anything"}],
            [],
        )["nodes"]
        assert nodes[0]["family"] == "Anything"

    def test_the_exported_page_filters_rather_than_recolours(self):
        """A node-link layout is an all-pairs context and the palette only
        clears its separation floor for three hues -- see _category."""
        src = pathlib.Path("src/export.py").read_text(encoding="utf-8")
        assert 'id="family-filter"' in src
        assert "familyFilter.addEventListener" in src

    def test_the_frontier_nodes_carry_no_family(self):
        from src.export import build_graph_payload

        payload = build_graph_payload(
            [], [{"source_channel_id": "s1", "target_channel_ref": "@ref",
                  "edge_type": "featured"}],
        )
        assert all(n["family"] == "" for n in payload["nodes"])


class TestFutureRunsGetThisWithoutBeingAsked:
    def test_the_node_runs_in_every_graph_run(self):
        src = pathlib.Path("src/graph.py").read_text(encoding="utf-8")
        assert 'graph.add_node("assign_niche_families"' in src
        assert 'graph.add_edge("populate_shared_fields", "assign_niche_families")' in src

    def test_the_gate_can_heal_it_if_the_run_could_not(self):
        from src.tools.export_completeness import _nodes

        assert "assign_niche_families" in _nodes()

    def test_it_heals_in_the_order_the_graph_runs_it(self):
        from src.tools.export_completeness import _nodes

        names = list(_nodes())
        assert names.index("assign_niche_families") < names.index("assign_cohorts")
